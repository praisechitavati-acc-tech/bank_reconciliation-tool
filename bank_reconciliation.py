import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook


BASE_DIR = Path(__file__).resolve().parent
BANK_STATEMENT_FILE = BASE_DIR / "Bank_Statement.csv"
CASHBOOK_FILE = BASE_DIR / "Cashbook.xlsx"


def create_sample_files():
    """Create sample input files when they are not already present."""
    if not BANK_STATEMENT_FILE.exists():
        bank_rows = [
            ["Date", "Description", "Amount"],
            ["2026-09-01", "Opening balance deposit", 1200.00],
            ["2026-09-02", "Office supplies", -85.50],
            ["2026-09-03", "Client payment", 450.00],
            ["2026-09-04", "Bank service fee", -25.00],
        ]

        with BANK_STATEMENT_FILE.open("w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            writer.writerows(bank_rows)

    if not CASHBOOK_FILE.exists():
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Cashbook"
        worksheet.append(["Date", "Amount", "Description"])
        worksheet.append(["2026-09-01", 1200.00, "Opening balance deposit"])
        worksheet.append(["2026-09-02", -85.50, "Office supplies"])
        worksheet.append(["2026-09-03", 450.00, "Client payment"])
        worksheet.append(["2026-09-05", -60.00, "Petty cash purchase"])
        workbook.save(CASHBOOK_FILE)


def read_bank_statement():
    transactions = []
    with BANK_STATEMENT_FILE.open("r", newline="", encoding="utf-8") as file:
        for row in csv.DictReader(file):
            transactions.append(
                {
                    "date": row["Date"],
                    "description": row["Description"],
                    "amount": float(row["Amount"]),
                }
            )
    return transactions


def read_cashbook():
    workbook = load_workbook(CASHBOOK_FILE, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(min_row=2, values_only=True)

    transactions = []
    for date, amount, description in rows:
        transactions.append(
            {
                "date": str(date),
                "description": str(description),
                "amount": float(amount),
            }
        )
    return transactions


def format_transaction(transaction):
    return (
        f"{transaction['date']} | "
        f"{transaction['description']} | "
        f"{transaction['amount']:,.2f}"
    )


def reconcile_transactions(bank_transactions, cashbook_transactions):
    """Compare transactions by amount and return the three reconciliation groups."""
    cashbook_by_amount = {
        transaction["amount"]: transaction for transaction in cashbook_transactions
    }
    bank_by_amount = {
        transaction["amount"]: transaction for transaction in bank_transactions
    }

    matched = [
        (bank_transaction, cashbook_by_amount[bank_transaction["amount"]])
        for bank_transaction in bank_transactions
        if bank_transaction["amount"] in cashbook_by_amount
    ]
    bank_only = [
        transaction
        for transaction in bank_transactions
        if transaction["amount"] not in cashbook_by_amount
    ]
    cashbook_only = [
        transaction
        for transaction in cashbook_transactions
        if transaction["amount"] not in bank_by_amount
    ]

    return matched, bank_only, cashbook_only


def print_reconciliation(matched, bank_only, cashbook_only):
    print("MATCHED TRANSACTIONS")
    print("--------------------")
    for bank_transaction, cashbook_transaction in matched:
        print(
            f"Bank:     {format_transaction(bank_transaction)}\n"
            f"Cashbook: {format_transaction(cashbook_transaction)}"
        )
    if not matched:
        print("None")

    print("\nIN BANK ONLY - MISSING IN CASHBOOK")
    print("----------------------------------")
    for transaction in bank_only:
        print(format_transaction(transaction))
    if not bank_only:
        print("None")

    print("\nIN CASHBOOK ONLY - MISSING IN BANK")
    print("----------------------------------")
    for transaction in cashbook_only:
        print(format_transaction(transaction))
    if not cashbook_only:
        print("None")


def main():
    create_sample_files()
    bank_transactions = read_bank_statement()
    cashbook_transactions = read_cashbook()
    matched, bank_only, cashbook_only = reconcile_transactions(
        bank_transactions, cashbook_transactions
    )
    print_reconciliation(matched, bank_only, cashbook_only)


if __name__ == "__main__":
    main()